"""Excel (.xlsx) and PDF export builders for a household's budget + bouwdepot + projects."""
import io


def _eur(v):
    s = f"{float(v or 0):,.2f}"
    return "€ " + s.replace(",", "X").replace(".", ",").replace("X", ".")


def build_excel(p):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    wb = Workbook()
    hh = p["household"]
    dash = p["dashboard"]
    persons = {pp["person_id"]: pp["name"] for pp in dash.get("persons", [])}
    head = Font(bold=True, color="FFFFFF")
    fill = PatternFill("solid", fgColor="0F172A")
    first = {"v": True}

    def sheet(title, headers, rows):
        ws = wb.active if first["v"] else wb.create_sheet(title)
        first["v"] = False
        ws.title = title
        ws.append(headers)
        for c in ws[1]:
            c.font = head
            c.fill = fill
        for r in rows:
            ws.append(r)
        for col in ws.columns:
            width = max((len(str(c.value)) for c in col if c.value is not None), default=10)
            ws.column_dimensions[col[0].column_letter].width = min(width + 3, 42)

    sheet("Dashboard",
          ["Maand", "Status", "Inkomen", "Vaste lasten", "Variabele lasten",
           "Totale lasten", "Over", "Cumulatief", "Spaarquote %"],
          [[m["label"], m["status"], m["income"], m["fixed"], m["variable"],
            m["expenses"], m["over"], m["cumulative"], m["savings_rate"]] for m in dash["months"]])
    a = dash["annual"]
    sheet("Jaartotaal", ["Post", "Bedrag"],
          [["Totaal inkomen", a["income"]], ["Vaste lasten", a["fixed"]],
           ["Variabele lasten", a["variable"]], ["Totale lasten", a["expenses"]],
           ["Over", a["over"]], ["Gem. per maand sparen", a.get("avg_monthly_over", 0)],
           ["Spaarquote %", a["savings_rate"]]])
    sheet("Per persoon", ["Persoon", "Verdient", "Betaalt", "Aandeel gezamenlijk", "Houdt over"],
          [[persons.get(pid, pid), v["income"], v["own"], v["joint_share"], v["net"]]
           for pid, v in dash["per_person_year"].items()])
    sheet("Categorieen", ["Categorie", "Uitgaven (jaar)"],
          [[k, v] for k, v in dash.get("expense_by_category", {}).items()])

    depot_rows, post_rows = [], []
    for d in p["depots"]:
        depot_rows.append([d["name"], d["start_amount"], d["paid_out"], d["submitted_not_paid"],
                           d["still_to_submit"], d["freely_available"], d["end_date"], d["days_remaining"]])
        for post in d["posts"]:
            post_rows.append([d["name"], post["name"], post["budget"], post["accepted_quotes"],
                              post["invoiced"], post["paid"], post["commitment"], post["room"]])
    sheet("Bouwdepot",
          ["Depot", "Startbedrag", "Uitbetaald", "Ingediend n.n. betaald", "Nog in te dienen",
           "Vrij besteedbaar", "Einddatum", "Dagen resterend"], depot_rows)
    sheet("Bouwposten",
          ["Depot", "Bouwpost", "Budget", "Geaccepteerde offertes", "Gefactureerd",
           "Betaald", "Verplichting", "Ruimte"], post_rows)
    sheet("Facturen",
          ["Leverancier", "Type", "Status", "Bedrag incl btw", "Factuurbedrag", "Ingediend", "Betaald op"],
          [[i.get("supplier"), i.get("type"), i.get("status"), i.get("amount_incl_vat"),
            i.get("invoice_amount"), "Ja" if i.get("submitted_to_bank") else "Nee", i.get("paid_on")]
           for i in p.get("invoices", [])])
    sheet("Doelen & Sparen",
          ["Doel", "Type", "Streefdatum", "Saldo", "Per maand", "Nog nodig", "Maanden", "Haalbaar"],
          [[g["name"], "Project" if g.get("has_target") else "Doorlopend",
            g.get("target_date") or "-", g["balance"], g["monthly_amount"],
            g.get("remaining", "-") if g.get("has_target") else "-",
            g.get("months_left", "-") if g.get("has_target") else "-",
            ("Ja" if g.get("feasible") else "Nee") if g.get("has_target") else "-"]
           for g in p.get("goals", [])])

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def build_pdf(p):
    from reportlab.lib.pagesizes import A4
    from reportlab.lib import colors
    from reportlab.lib.units import mm
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle

    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, topMargin=18 * mm, bottomMargin=15 * mm,
                            leftMargin=16 * mm, rightMargin=16 * mm)
    styles = getSampleStyleSheet()
    title = ParagraphStyle('t', parent=styles['Title'], textColor=colors.HexColor('#0f172a'))
    h2 = ParagraphStyle('h2', parent=styles['Heading2'], textColor=colors.HexColor('#0f172a'))
    hh = p["household"]
    dash = p["dashboard"]
    a = dash["annual"]
    persons = {pp["person_id"]: pp["name"] for pp in dash.get("persons", [])}
    el = [Paragraph(f"Huishoudbudget &amp; Bouwdepot — {hh['name']}", title),
          Paragraph(f"Jaaroverzicht {dash['year']} · valuta EUR", styles['Normal']),
          Spacer(1, 8)]

    def tbl(data):
        t = Table(data, hAlign='LEFT')
        t.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#0f172a')),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
            ('FONTSIZE', (0, 0), (-1, -1), 8),
            ('GRID', (0, 0), (-1, -1), 0.3, colors.HexColor('#cbd5e1')),
            ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#f1f5f9')]),
            ('ALIGN', (1, 0), (-1, -1), 'RIGHT'),
            ('LEFTPADDING', (0, 0), (-1, -1), 6), ('RIGHTPADDING', (0, 0), (-1, -1), 6),
        ]))
        return t

    el.append(Paragraph("Jaartotaal", h2))
    el.append(tbl([["Post", "Bedrag"],
                   ["Totaal inkomen", _eur(a['income'])],
                   ["Vaste lasten", _eur(a['fixed'])],
                   ["Variabele lasten", _eur(a['variable'])],
                   ["Totale lasten", _eur(a['expenses'])],
                   ["Over", _eur(a['over'])],
                   ["Gem. per maand sparen", _eur(a.get('avg_monthly_over', 0))],
                   ["Spaarquote", f"{a['savings_rate']}%"]]))
    el.append(Spacer(1, 10))
    el.append(Paragraph("Per persoon — wat houden we over?", h2))
    el.append(tbl([["Persoon", "Verdient", "Betaalt", "Gezamenlijk", "Houdt over"]] +
                  [[persons.get(pid, pid), _eur(v['income']), _eur(v['own']),
                    _eur(v['joint_share']), _eur(v['net'])]
                   for pid, v in dash['per_person_year'].items()]))

    for d in p["depots"]:
        el.append(Spacer(1, 10))
        el.append(Paragraph(f"Bouwdepot — {d['name']}", h2))
        el.append(tbl([["Post", "Bedrag"],
                       ["Startbedrag", _eur(d['start_amount'])],
                       ["Uitbetaald", _eur(d['paid_out'])],
                       ["Ingediend, nog niet betaald", _eur(d['submitted_not_paid'])],
                       ["Nog in te dienen", _eur(d['still_to_submit'])],
                       ["Vrij besteedbaar", _eur(d['freely_available'])],
                       ["Einddatum", str(d['end_date'] or '-')],
                       ["Dagen resterend", str(d['days_remaining'] if d['days_remaining'] is not None else '-')]]))
        if d["posts"]:
            el.append(Spacer(1, 4))
            el.append(tbl([["Bouwpost", "Budget", "Verplichting", "Ruimte"]] +
                          [[po['name'], _eur(po['budget']), _eur(po['commitment']), _eur(po['room'])]
                           for po in d['posts']]))

    if p.get("goals"):
        el.append(Spacer(1, 10))
        el.append(Paragraph("Doelen &amp; Sparen", h2))
        el.append(tbl([["Doel", "Type", "Streefdatum", "Saldo", "Per maand", "Haalbaar"]] +
                      [[g['name'], "Project" if g.get('has_target') else "Doorlopend",
                        str(g.get('target_date') or '-'), _eur(g['balance']),
                        _eur(g.get('required_monthly') if g.get('has_target') else g['monthly_amount']),
                        ('Ja' if g.get('feasible') else 'Nee') if g.get('has_target') else '-']
                       for g in p['goals']]))

    doc.build(el)
    return buf.getvalue()
